import argparse
import calendar
import random
import secrets
from datetime import date
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError as error:
    raise SystemExit("Install the required package with: python -m pip install openpyxl") from error


SHEET_NAME = "Danh sách phiếu điều tra"
MEMBERS_PER_HOUSEHOLD = 4

ETHNICITIES = ["Kinh", "Tày", "Thái", "Mường", "Nùng", "Dao", "Hmong"]
FAMILY_NAMES = ["Nguyễn", "Trần", "Lê", "Phạm", "Hoàng", "Huỳnh", "Phan", "Vũ", "Võ", "Đặng"]
MIDDLE_NAMES = ["Văn", "Thị", "Hữu", "Ngọc", "Minh", "Gia", "Đức", "Thu", "Thanh", "Quốc"]
GIVEN_NAMES = [
    "An", "Bình", "Chi", "Dũng", "Hà", "Hạnh", "Huy", "Khánh", "Linh", "Long",
    "Mai", "Nam", "Ngọc", "Phúc", "Quân", "Sơn", "Thảo", "Trang", "Tuấn", "Vy",
]

GROUP_HEADERS = [
    ("B1", "K1", "Thông tin đối tượng"),
    ("L1", "Q1", "Thông tin nơi đang ở"),
    ("R1", "AF1", "Thông tin học tập"),
    ("AG1", "AI1", "Xóa mù chữ"),
    ("AJ1", "AS1", "Khuyết tật"),
    ("AT1", "AU1", "Hoàn cảnh"),
    ("AV1", "AY1", "Thông tin liên lạc"),
]

HEADERS = {
    "B2": "TT", "C2": "Họ đệm", "D2": "Tên", "E2": "Ngày", "F2": "Tháng",
    "G2": "Năm sinh", "H2": "Nữ", "I2": "Dân tộc", "J2": "Tôn giáo",
    "K2": "Diện ưu tiên", "L2": "Chủ hộ", "L3": "Họ đệm", "M3": "Tên",
    "N2": "Địa chỉ: số nhà, tên đường, tổ (nếu có)", "O2": "Số phiếu",
    "P2": "Diện cư trú", "Q2": "Tình trạng cư trú", "R2": "Khối học",
    "S2": "Lớp học", "V2": "Mã trường", "W2": "Bậc tốt nghiệp", "X2": "Bổ túc",
    "Y2": "Năm tốt nghiệp", "Z2": "Bậc TN nghề", "AB2": "Năm TN nghề",
    "AC2": "Học xong", "AC3": "Lớp", "AD3": "Năm", "AE2": "Bỏ học",
    "AE3": "Lớp", "AF3": "Năm", "AG2": "Đang học lớp", "AH2": "Hoàn thành lớp",
    "AI2": "Tái mù chữ mức", "AJ2": "Khuyết tật vận động", "AK2": "Khuyết tật nghe nói",
    "AL2": "Khuyết tật nhìn", "AM2": "Khuyết tật thần kinh, tâm thần",
    "AN2": "Khuyết tật trí tuệ", "AO2": "Khuyết tật học tập", "AP2": "Tự kỷ",
    "AQ2": "Khuyết tật khác", "AR2": "Có chứng nhận khuyết tật",
    "AS2": "Khả năng học tập", "AT2": "Hoàn cảnh đặc biệt",
    "AU2": "Chi tiết hoàn cảnh đặt biệt", "AV2": "Quan hệ với chủ hộ",
    "AW2": "Họ tên cha hoặc mẹ", "AX2": "Điện thoại", "AY2": "Ghi chú",
}

HEADER_MERGES = [
    "B2:B3", "C2:C3", "D2:D3", "E2:E3", "F2:F3", "G2:G3", "H2:H3", "I2:I3",
    "J2:J3", "K2:K3", "L2:M2", "N2:N3", "O2:O3", "P2:P3", "Q2:Q3", "R2:R3",
    "S2:S3", "V2:V3", "W2:W3", "X2:X3", "Y2:Y3", "Z2:Z3", "AB2:AB3",
    "AC2:AD2", "AE2:AF2", "AG2:AG3", "AH2:AH3", "AI2:AI3", "AJ2:AJ3",
    "AK2:AK3", "AL2:AL3", "AM2:AM3", "AN2:AN3", "AO2:AO3", "AP2:AP3",
    "AQ2:AQ3", "AR2:AR3", "AS2:AS3", "AT2:AT3", "AU2:AU3", "AV2:AV3",
    "AW2:AW3", "AX2:AX3", "AY2:AY3",
]


def make_name(rng):
    return f"{rng.choice(FAMILY_NAMES)} {rng.choice(MIDDLE_NAMES)}", rng.choice(GIVEN_NAMES)


def make_unique_name(rng, used_names):
    for _ in range(100):
        name = make_name(rng)
        if name not in used_names:
            used_names.add(name)
            return name
    raise RuntimeError("Could not generate a unique name for the household.")


def make_birth_date(rng, age):
    current_year = date.today().year
    year = current_year - age
    month = rng.randint(1, 12)
    day = rng.randint(1, calendar.monthrange(year, month)[1])
    return day, month, year


def style_template(sheet):
    group_fill = PatternFill("solid", fgColor="276749")
    header_fill = PatternFill("solid", fgColor="DCE6F1")
    index_fill = PatternFill("solid", fgColor="EEF2F7")
    header_font = Font(name="Arial", size=9, bold=True, color="1F2937")
    white_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    thin_gray = Side(style="thin", color="AAB4C0")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    for start, end, title in GROUP_HEADERS:
        sheet.merge_cells(f"{start}:{end}")
        cell = sheet[start]
        cell.value = title
        cell.fill = group_fill
        cell.font = white_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for cell_ref, title in HEADERS.items():
        cell = sheet[cell_ref]
        cell.value = title
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    sequence = 1
    for column_index in range(2, 52):
        column = get_column_letter(column_index)
        if column in {"T", "U", "AA"}:
            continue
        cell = sheet.cell(row=4, column=column_index, value=sequence)
        cell.fill = index_fill
        cell.font = Font(name="Arial", size=8, color="526071")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
        sequence += 1

    for merge_range in HEADER_MERGES:
        sheet.merge_cells(merge_range)

    widths = {
        "B": 5, "C": 18, "D": 14, "E": 7, "F": 7, "G": 9, "H": 6, "I": 14,
        "J": 12, "K": 14, "L": 16, "M": 12, "N": 25, "O": 25, "P": 14,
        "Q": 18, "R": 10, "S": 12, "V": 18, "W": 16, "X": 9, "Y": 12,
        "Z": 14, "AB": 13, "AC": 9, "AD": 9, "AE": 9, "AF": 9, "AG": 14,
        "AH": 14, "AI": 14, "AJ": 15, "AK": 15, "AL": 13, "AM": 20,
        "AN": 15, "AO": 15, "AP": 10, "AQ": 15, "AR": 17, "AS": 15,
        "AT": 18, "AU": 24, "AV": 18, "AW": 20, "AX": 15, "AY": 24,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    sheet.row_dimensions[1].height = 32
    sheet.row_dimensions[2].height = 44
    sheet.row_dimensions[3].height = 24
    sheet.row_dimensions[4].height = 20
    sheet.freeze_panes = "C5"


def write_sample_rows(sheet, row_count, rng):
    timestamp = date.today().strftime("%Y%m%d")
    token = secrets.token_hex(2).upper()
    sequence = 1
    household_number = 1
    row = 5

    while sequence <= row_count:
        household_code = f"DEMO_{timestamp}_{token}_{household_number:04d}"
        head_last_name, head_first_name = make_name(rng)
        used_names = {(head_last_name, head_first_name)}
        head_gender = rng.choice(["Nam", "Nữ"])
        member_count = min(MEMBERS_PER_HOUSEHOLD, row_count - sequence + 1)

        for member_index in range(member_count):
            if member_index == 0:
                last_name, first_name = head_last_name, head_first_name
                gender = head_gender
                relationship = "Chủ hộ"
                age = rng.randint(30, 65)
            elif member_index == 1:
                last_name, first_name = make_unique_name(rng, used_names)
                gender = "Nữ" if head_gender == "Nam" else "Nam"
                relationship = "Vợ" if head_gender == "Nam" else "Chồng"
                age = rng.randint(27, 62)
            else:
                last_name, first_name = make_unique_name(rng, used_names)
                gender = rng.choice(["Nam", "Nữ"])
                relationship = "Con"
                age = rng.randint(3, 22)

            day, month, year = make_birth_date(rng, age)
            values = {
                "B": sequence,
                "C": last_name,
                "D": first_name,
                "E": day,
                "F": month,
                "G": year,
                "H": "X" if gender == "Nữ" else None,
                "I": rng.choice(ETHNICITIES),
                "L": head_last_name,
                "M": head_first_name,
                "N": f"Số {rng.randint(1, 250)}, đường Mẫu, tổ {rng.randint(1, 12)}",
                "O": household_code,
                "P": "Thường trú",
                "Q": "Đang ở",
                "AV": relationship,
                "AW": f"{head_last_name} {head_first_name}" if relationship == "Con" else None,
                "AX": f"09{rng.randint(10000000, 99999999)}" if member_index == 0 else None,
                "AY": "Dữ liệu mẫu tạo tự động",
            }
            for column, value in values.items():
                if value is not None:
                    cell = sheet[f"{column}{row}"]
                    cell.value = value
                    cell.alignment = Alignment(vertical="center")
                    if column in {"B", "E", "F", "G", "H"}:
                        cell.alignment = Alignment(horizontal="center", vertical="center")
            row += 1
            sequence += 1

        household_number += 1

    last_row = row - 1
    sheet.auto_filter.ref = f"B4:AY{last_row}"
    sheet.row_dimensions[4].height = 20
    return last_row, household_number - 1


def main():
    parser = argparse.ArgumentParser(description="Generate import-ready sample household data.")
    parser.add_argument("--rows", type=int, default=1000, help="Number of person rows to generate (default: 1000).")
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(__file__).resolve().parent / "sample_import_1000.xlsx",
        help="Output .xlsx path.",
    )
    parser.add_argument("--seed", type=int, help="Optional random seed for repeatable names and dates.")
    args = parser.parse_args()

    if args.rows < 1:
        parser.error("--rows must be greater than zero")
    if args.output.suffix.lower() != ".xlsx":
        parser.error("--output must use the .xlsx extension")

    rng = random.Random(args.seed)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    style_template(sheet)
    last_row, household_count = write_sample_rows(sheet, args.rows, rng)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    print(f"Created: {args.output}")
    print(f"Person rows: {args.rows}; households: {household_count}; data rows: 5-{last_row}")
    print("Upload this XLSX with the same school year selected in the import form.")


if __name__ == "__main__":
    main()